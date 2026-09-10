import io
import json
import os
import re
from typing import Dict, List, Tuple

import pymupdf
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


SKILL_LIBRARY = {
    # Programming
    "python", "java", "c", "c++", "c#", "javascript", "typescript",
    "html", "css", "sql", "r", "php", "kotlin", "swift",
    # Data / ML
    "machine learning", "deep learning", "artificial intelligence",
    "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch",
    "nlp", "computer vision", "data analysis", "data visualization",
    "statistics", "power bi", "tableau", "excel",
    # Backend / Web
    "fastapi", "flask", "django", "react", "node.js", "rest api",
    "api", "microservices",
    # Databases
    "mysql", "postgresql", "mongodb", "sqlite", "redis",
    # Cloud / DevOps
    "aws", "azure", "gcp", "docker", "kubernetes", "git", "github",
    "ci/cd", "linux",
    # Tools / soft skills
    "jira", "agile", "scrum", "communication", "leadership",
    "problem solving", "teamwork",
}

SECTION_ALIASES = {
    "Summary / Objective": ["summary", "objective", "profile", "professional summary"],
    "Experience": ["experience", "work experience", "employment"],
    "Education": ["education", "academic background", "qualifications"],
    "Skills": ["skills", "technical skills", "core skills", "technologies"],
    "Projects": ["projects", "academic projects", "personal projects"],
    "Certifications": ["certifications", "certificates"],
}

ACTION_VERBS = {
    "built", "developed", "implemented", "designed", "created", "automated",
    "optimized", "improved", "analyzed", "deployed", "led", "managed",
    "integrated", "engineered", "reduced", "increased", "delivered",
}


def normalize(text: str) -> str:
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_pdf(file_bytes: bytes) -> str:
    chunks = []
    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            chunks.append(page.get_text("text"))
    return "\n".join(chunks)


def extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs + table_text)


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")


def detect_sections(text: str) -> Dict[str, bool]:
    lower = normalize(text)
    return {
        section: any(alias in lower for alias in aliases)
        for section, aliases in SECTION_ALIASES.items()
    }


def detect_skills(text: str) -> List[str]:
    lower = normalize(text)
    found = []
    for skill in SKILL_LIBRARY:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    return sorted(found)


def extract_job_skills(job_description: str) -> List[str]:
    if not job_description:
        return []
    return detect_skills(job_description)


def keyword_coverage(resume_text: str, job_description: str) -> Tuple[int, List[str], List[str]]:
    if not job_description.strip():
        return 0, [], []

    resume_skills = set(detect_skills(resume_text))
    job_skills = set(extract_job_skills(job_description))

    if not job_skills:
        return 0, [], []

    matched = sorted(resume_skills & job_skills)
    missing = sorted(job_skills - resume_skills)
    score = round((len(matched) / len(job_skills)) * 100)
    return score, matched, missing


def semantic_job_match(resume_text: str, job_description: str) -> int:
    if not job_description.strip():
        return 0

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            max_features=5000,
        )
        matrix = vectorizer.fit_transform([resume_text, job_description])
        sim = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]
        return int(round(sim * 100))
    except ValueError:
        return 0


def ats_score(text: str, sections: Dict[str, bool]) -> Tuple[int, List[str], List[str]]:
    issues = []
    strengths = []
    score = 100

    clean = text.strip()

    if len(clean) < 500:
        score -= 20
        issues.append("Resume text is very short; add more evidence of skills, projects, and achievements.")

    if len(clean) > 7000:
        score -= 10
        issues.append("Resume is quite long. Consider reducing repetition and keeping the content focused.")

    if not sections["Skills"]:
        score -= 12
        issues.append("No clear Skills section detected.")

    if not sections["Education"]:
        score -= 8
        issues.append("No clear Education section detected.")

    if not sections["Experience"] and not sections["Projects"]:
        score -= 12
        issues.append("No clear Experience or Projects section detected.")

    if not sections["Summary / Objective"]:
        score -= 5
        issues.append("Consider adding a concise professional summary.")

    lines = [x.strip() for x in text.splitlines() if x.strip()]
    bullet_lines = sum(1 for x in lines if re.match(r"^[-•*]", x))
    if len(lines) >= 10 and bullet_lines >= 3:
        strengths.append("Uses bullet-style content, which is easier to scan.")
    else:
        score -= 5
        issues.append("Use more concise bullet points for responsibilities and achievements.")

    quantified = len(re.findall(r"\b\d+%|\b\d+\+|\b\d{2,}\b", text))
    if quantified >= 3:
        strengths.append("Contains measurable numbers/metrics.")
    else:
        score -= 7
        issues.append("Add measurable outcomes where possible, such as %, time saved, users, accuracy, or revenue.")

    lower = normalize(text)
    action_count = sum(1 for verb in ACTION_VERBS if re.search(r"\b" + re.escape(verb) + r"\b", lower))
    if action_count >= 4:
        strengths.append("Uses several strong action verbs.")
    else:
        score -= 5
        issues.append("Start more bullets with action verbs such as built, implemented, optimized, or analyzed.")

    # Penalize extremely unusual character density often caused by extraction problems.
    alpha = sum(c.isalpha() for c in clean)
    strange = sum(not (c.isalnum() or c.isspace() or c in ".,;:!?@#$%&()[]{}+-_/|•'\"") for c in clean)
    if alpha > 0 and strange / max(len(clean), 1) > 0.08:
        score -= 5
        issues.append("Some extracted characters look unusual; check the original PDF formatting.")

    return max(0, min(100, score)), issues, strengths


def analyze_resume(uploaded_file, job_description: str = "") -> Dict:
    result = {
        "filename": uploaded_file.name,
        "errors": [],
        "text": "",
        "sections": {},
        "skills_found": [],
        "missing_skills": [],
        "matched_skills": [],
        "ats_score": 0,
        "keyword_score": 0,
        "job_match_score": 0,
        "overall_score": 0,
        "issues": [],
        "strengths": [],
    }

    try:
        text = extract_text(uploaded_file)
    except Exception as exc:
        result["errors"].append(f"Could not read file: {exc}")
        result["json"] = json.dumps(result, indent=2)
        return result

    if not text.strip():
        result["errors"].append(
            "No readable text was found. Scanned/image-only PDFs need OCR support."
        )
        result["json"] = json.dumps(result, indent=2)
        return result

    result["text"] = text
    result["sections"] = detect_sections(text)
    result["skills_found"] = detect_skills(text)

    ats, issues, strengths = ats_score(text, result["sections"])
    result["ats_score"] = ats
    result["issues"].extend(issues)
    result["strengths"].extend(strengths)

    kw_score, matched, missing = keyword_coverage(text, job_description)
    result["keyword_score"] = kw_score
    result["matched_skills"] = matched
    result["missing_skills"] = missing

    semantic = semantic_job_match(text, job_description)
    # With a JD: combine skill overlap and semantic similarity.
    # Without a JD: use ATS score only, avoiding a misleading job-match score.
    if job_description.strip():
        result["job_match_score"] = round(0.6 * kw_score + 0.4 * semantic)
        result["overall_score"] = round(0.55 * ats + 0.45 * result["job_match_score"])
    else:
        result["job_match_score"] = 0
        result["overall_score"] = ats

    result["json"] = json.dumps(result, indent=2)
    return result


def generate_ai_feedback(resume_text: str, job_description: str, result: dict):
    api_key = None
    try:
        import streamlit as st
        api_key = st.secrets.get("OPENAI_API_KEY")
    except Exception:
        api_key = os.getenv("OPENAI_API_KEY")

    if not api_key:
        return "", "OPENAI_API_KEY is not configured. Add it to .streamlit/secrets.toml locally or Streamlit Secrets in production."

    try:
        from openai import OpenAI

        client = OpenAI(api_key=api_key)

        prompt = f"""
You are an ATS resume coach.

Analyze the resume below. Do not invent experience, education, skills, or achievements.
Give practical, truthful improvements.

RESUME:
{resume_text[:12000]}

JOB DESCRIPTION:
{job_description[:10000] if job_description else "(not provided)"}

RULE-BASED RESULTS:
{json.dumps({
    "overall_score": result["overall_score"],
    "ats_score": result["ats_score"],
    "job_match_score": result["job_match_score"],
    "skills_found": result["skills_found"],
    "matched_skills": result["matched_skills"],
    "missing_skills": result["missing_skills"],
    "issues": result["issues"],
}, indent=2)}

Return markdown with:
1. Top 5 improvements
2. Missing/weak keywords
3. Three rewritten bullet examples based ONLY on information already present
4. A stronger 3-4 sentence professional summary, without inventing facts
5. ATS formatting advice
"""

        response = client.responses.create(
            model="gpt-5",
            input=prompt,
        )
        return response.output_text, None

    except Exception as exc:
        return "", f"AI feedback failed: {exc}"
